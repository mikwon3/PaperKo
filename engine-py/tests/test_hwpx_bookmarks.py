"""The .hwpx we write carries no bookmarks.

The skeleton (``hwpx/assets/default_skeleton.hwpx``) is a real Hancom document, so it
holds that document's 88 bookmarks — with its names, like ``asad-v1.2-사용자-매뉴얼``.
The converter copies skeleton paragraphs as templates, so those bookmarks rode along
into every exported file and showed up in Hangul's bookmark list (2026-10-29, reported
by a user; a two-paragraph document carried two of them).

We also check the fields stay paired: a ``fieldBegin`` without its ``fieldEnd`` makes
Hangul declare the document damaged and "repair" it on open.
"""

import re
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document as Docx

ENGINE = Path(__file__).resolve().parents[1] / "translate_engine"
SCRIPTS = ENGINE / "hwpx" / "scripts"
SKELETON = ENGINE / "hwpx" / "assets" / "default_skeleton.hwpx"


def _converter():
    """Import the bundled converter the way docx_export does."""
    if str(SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SCRIPTS))
    import importlib

    importlib.import_module("omml_to_hwp")  # sibling import used by convert.py
    return importlib.import_module("convert")


def _sample_docx(path: Path) -> Path:
    doc = Docx()
    doc.add_heading("책갈피 확인", level=1)
    doc.add_paragraph("첫 문단입니다.")
    doc.add_heading("작은 제목", level=2)
    doc.add_paragraph("글머리표 하나", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "가"
    table.cell(0, 1).text = "나"
    doc.save(path)
    return path


@pytest.mark.skipif(not SKELETON.exists(), reason="bundled skeleton missing")
def test_exported_hwpx_has_no_bookmarks(tmp_path):
    conv = _converter()
    docx = _sample_docx(tmp_path / "in.docx")
    out = tmp_path / "out.hwpx"
    conv.convert(str(docx), str(SKELETON), str(out), verbose=False)

    with zipfile.ZipFile(out) as z:
        section = z.read("Contents/section0.xml").decode("utf-8")

    names = re.findall(r'type="BOOKMARK"\s+name="([^"]*)"', section)
    assert names == [], f"bookmarks rode along from the skeleton: {names}"

    begins = len(re.findall(r"<hp:fieldBegin", section))
    ends = len(re.findall(r"<hp:fieldEnd", section))
    assert begins == ends, f"unpaired fields: {begins} begins, {ends} ends"

    # Stripping must not take the body with it.
    for text in ("첫 문단입니다.", "작은 제목", "글머리표 하나"):
        assert text in section, f"body text went missing: {text!r}"


@pytest.mark.skipif(not SKELETON.exists(), reason="bundled skeleton missing")
def test_skeleton_still_has_the_bookmarks_we_strip():
    """Guards the premise: if the skeleton is ever replaced by a clean one, the test
    above would pass for the wrong reason and stop protecting anything."""
    with zipfile.ZipFile(SKELETON) as z:
        section = z.read("Contents/section0.xml").decode("utf-8")
    assert 'type="BOOKMARK"' in section
